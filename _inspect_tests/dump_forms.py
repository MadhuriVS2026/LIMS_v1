"""
Inspection utility: dump the structure of every R&D test-type Excel form so we
can understand which fields are inputs, which are auto-calculated (formulas),
and which come from the chromatography system (Area fields).

Not part of the application — analysis tooling only.
"""
import sys
from pathlib import Path

FORMS_DIR = Path(r"d:\Anti Gravity LIMS\rd-lab-instance\Test types")
OUT = Path(__file__).parent / "forms_dump.txt"


def dump_xlsx(path: Path, out) -> None:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:
        out.write(f"  !! could not open: {exc}\n")
        return

    for ws in wb.worksheets:
        out.write(f"  --- sheet: {ws.title} (dims={ws.dimensions}, max_row={ws.max_row}, max_col={ws.max_column}) ---\n")
        max_r = min(ws.max_row or 0, 80)
        max_c = min(ws.max_column or 0, 20)
        for r in range(1, max_r + 1):
            row_cells = []
            for c in range(1, max_c + 1):
                cell = ws.cell(row=r, column=c)
                v = cell.value
                if v is None or (isinstance(v, str) and not v.strip()):
                    continue
                kind = "F" if isinstance(v, str) and v.startswith("=") else "V"
                row_cells.append(f"{cell.coordinate}[{kind}]={v!r}")
            if row_cells:
                out.write("    " + " | ".join(row_cells) + "\n")


def dump_xls(path: Path, out) -> None:
    import xlrd

    try:
        wb = xlrd.open_workbook(str(path), formatting_info=False)
    except Exception as exc:
        out.write(f"  !! could not open: {exc}\n")
        return

    for ws in wb.sheets():
        out.write(f"  --- sheet: {ws.name} (rows={ws.nrows}, cols={ws.ncols}) ---\n")
        for r in range(min(ws.nrows, 80)):
            row_cells = []
            for c in range(min(ws.ncols, 20)):
                v = ws.cell_value(r, c)
                if v is None or (isinstance(v, str) and not v.strip()) or v == "":
                    continue
                row_cells.append(f"r{r + 1}c{c + 1}={v!r}")
            if row_cells:
                out.write("    " + " | ".join(row_cells) + "\n")


def main() -> None:
    files = sorted(FORMS_DIR.iterdir())
    only = sys.argv[1:] if len(sys.argv) > 1 else None

    with OUT.open("w", encoding="utf-8") as out:
        for f in files:
            if only and not any(o.lower() in f.name.lower() for o in only):
                continue
            out.write(f"\n{'=' * 100}\nFILE: {f.name}\n{'=' * 100}\n")
            if f.suffix.lower() == ".xlsx":
                dump_xlsx(f, out)
            elif f.suffix.lower() == ".xls":
                dump_xls(f, out)
            else:
                out.write("  (skipped: unsupported extension)\n")

    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
